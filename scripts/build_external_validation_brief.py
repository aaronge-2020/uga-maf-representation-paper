#!/usr/bin/env python3
"""Build the one-page Aaron-facing external-validation results brief."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "external_validation" / "aaron_email"
OUT = RESULTS / "Aaron_External_Validation_One_Page_Summary.docx"

INK = "172B4D"
BLUE = "2E74B5"
MUTED = "5F6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def set_run(run, size=10, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=100, bottom=65, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_table(table, header=True, font_size=8.3, first_col_left=True):
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            shade_cell(cell, LIGHT_GRAY if header and row_index == 0 else WHITE)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if first_col_left and col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run(run, size=font_size, bold=(header and row_index == 0), color=INK)


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)
    return paragraph


def add_rich_paragraph(doc, pieces, after=3, size=9.4, shade=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    for text, bold, color, italic in pieces:
        set_run(paragraph.add_run(text), size=size, bold=bold, color=color, italic=italic)
    if shade:
        p_pr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), shade)
        p_pr.append(shd)
        spacing = p_pr.find(qn("w:spacing"))
        spacing.set(qn("w:before"), "90")
        spacing.set(qn("w:after"), "90")
    return paragraph


def metric(frame, cohort, analysis, metric_name, cutoff=None):
    rows = frame[(frame["cohort"] == cohort) & (frame["analysis"] == analysis) & (frame["metric"] == metric_name)]
    if cutoff is not None:
        rows = rows[rows["cutoff"] == cutoff]
    if len(rows) != 1:
        raise RuntimeError(
            f"Expected one metric row for cohort={cohort}, analysis={analysis}, "
            f"metric={metric_name}, cutoff={cutoff}; found {len(rows)}"
        )
    return rows.iloc[0]


def pcawg_metric(frame, metric_name):
    rows = frame[frame["metric"].eq(metric_name)]
    if len(rows) != 1:
        raise RuntimeError(f"Expected one PCAWG metric row for metric={metric_name}; found {len(rows)}")
    return rows.iloc[0]


def parse_boolean_column(series):
    if series.dtype == bool:
        return series
    parsed = series.astype(str).str.strip().str.lower().map({"true": True, "false": False})
    if parsed.isna().any():
        raise RuntimeError("The POG570 included column contains values other than true/false")
    return parsed


def ci(row, digits=3):
    return f"{row['estimate']:.{digits}f} ({row['ci_low']:.{digits}f}-{row['ci_high']:.{digits}f})"


def build():
    pog = pd.read_csv(RESULTS / "pog570" / "pog570_primary_and_all_metrics_with_95ci.csv")
    pcawg = pd.read_csv(RESULTS / "pcawg" / "pcawg_metrics_with_95ci.csv")
    pog_predictions = pd.read_csv(RESULTS / "pog570" / "pog570_frozen_model_predictions.csv")
    pog_mapping = pd.read_csv(RESULTS / "pog570" / "pog570_sample_mapping_and_exclusions.csv")
    pcawg_predictions = pd.read_csv(RESULTS / "pcawg" / "pcawg_frozen_model_predictions.csv")
    pcawg_mapping = pd.read_csv(RESULTS / "pcawg" / "pcawg_913_sample_mapping.csv")
    pog_versions = json.loads((RESULTS / "pog570" / "software_versions_and_preprocessing.json").read_text())
    pcawg_versions = json.loads((RESULTS / "pcawg" / "software_versions.json").read_text())
    hrd_manifest = json.loads(
        (ROOT / "results" / "frozen_models" / "HRD_Score__signatures_plus_MAF_stack" / "manifest.json").read_text()
    )
    cancer_manifest = json.loads(
        (ROOT / "results" / "frozen_models" / "cancer_type_top20__standard_sbs96_id83" / "manifest.json").read_text()
    )

    pog_usable = int(parse_boolean_column(pog_mapping["included"]).sum())
    pog_excluded = int(len(pog_mapping) - pog_usable)
    breast_n = int(pog_predictions["analysis_group"].eq("primary_breast").sum())
    exploratory = pog_predictions[pog_predictions["analysis_group"].eq("exploratory_other")]
    exploratory_n = len(exploratory)
    exploratory_types = int(exploratory["CANCER_TYPE"].nunique())
    observed_mean = float(pog_predictions["HRD_SCORE"].mean())
    predicted_mean = float(pog_predictions["predicted_HRD_SCORE"].mean())
    breast_r2 = float(metric(pog, "primary_breast", "continuous", "r2")["estimate"])
    pcawg_usable = len(pcawg_predictions)
    pcawg_excluded = int(len(pcawg_mapping) - pcawg_usable)
    mapping_count = int(pcawg_mapping["true_label"].nunique())
    prepared_date = datetime.now().astimezone().strftime("%d %B %Y")
    commit = pog_versions.get("git_commit") or pcawg_versions.get("git_commit") or "unknown"
    hrd_feature_count = int(hrd_manifest["n_features"])
    cancer_feature_count = int(cancer_manifest["n_features"])

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run(header.add_run("UGA MAF Representation Paper  |  External Validation"), size=8, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    set_run(footer.add_run(f"Prepared {prepared_date}  |  Commit {commit[:7]}  |  Page 1"), size=7.5, color=MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(1)
    set_run(title.add_run("External Validation Results"), size=21, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(7)
    set_run(subtitle.add_run("POG570 HRD score and PCAWG ICGC-only cancer-type transfer"), size=10.5, color=MUTED)

    add_rich_paragraph(
        doc,
        [
            ("Takeaway. ", True, INK, False),
            ("POG570 retained useful HRD ranking but showed strong upward calibration shift. PCAWG transfer was heterogeneous across cancer types. No external data were used for fitting, tuning, cutoff selection, or post hoc adjustment.", False, INK, False),
        ],
        after=5,
        size=9.2,
        shade=LIGHT_BLUE,
    )

    add_heading(doc, "1. POG570 primary HRD validation")
    add_rich_paragraph(
        doc,
        [
            ("Cohort: ", True, INK, False),
            (f"{pog_usable}/{len(pog_mapping)} usable; {pog_excluded} exclusions. Primary analysis: {breast_n} breast cancers. Exploratory: {exploratory_n} non-breast cancers across {exploratory_types} types. ", False, INK, False),
            ("Model: ", True, INK, False),
            (f"frozen HRD_Score / signatures_plus_MAF_stack ({hrd_feature_count:,} features).", False, INK, False),
        ],
        after=3,
    )

    cont_table = doc.add_table(rows=1, cols=6)
    cont_table.style = "Table Grid"
    for index, text in enumerate(["Cohort", "N", "Pearson r", "Spearman rho", "MAE", "RMSE"]):
        cont_table.rows[0].cells[index].text = text
    for cohort_key, label, n in [("primary_breast", "Breast", breast_n), ("all_pog570", "All POG570", pog_usable)]:
        cells = cont_table.add_row().cells
        vals = [
            label,
            str(n),
            ci(metric(pog, cohort_key, "continuous", "pearson_r")),
            ci(metric(pog, cohort_key, "continuous", "spearman_rho")),
            ci(metric(pog, cohort_key, "continuous", "mae"), 2),
            ci(metric(pog, cohort_key, "continuous", "rmse"), 2),
        ]
        for cell, value in zip(cells, vals):
            cell.text = value
    set_table_geometry(cont_table, [1450, 560, 1860, 1860, 1815, 1815])
    format_table(cont_table, font_size=7.7)

    add_rich_paragraph(
        doc,
        [("Breast cancers at prespecified cutoffs", True, INK, False), ("  (estimate; 95% CI shown for balanced accuracy and AUROC)", False, MUTED, True)],
        after=2,
        size=8.5,
    )
    cutoff_table = doc.add_table(rows=1, cols=6)
    cutoff_table.style = "Table Grid"
    for index, text in enumerate(["Cutoff", "Accuracy", "Balanced accuracy", "Sensitivity", "Specificity", "AUROC"]):
        cutoff_table.rows[0].cells[index].text = text
    for cutoff in (24, 33, 42):
        cells = cutoff_table.add_row().cells
        vals = [
            str(cutoff),
            f"{metric(pog, 'primary_breast', 'fixed_cutoff', 'accuracy', cutoff)['estimate']:.3f}",
            ci(metric(pog, "primary_breast", "fixed_cutoff", "balanced_accuracy", cutoff)),
            f"{metric(pog, 'primary_breast', 'fixed_cutoff', 'sensitivity', cutoff)['estimate']:.3f}",
            f"{metric(pog, 'primary_breast', 'fixed_cutoff', 'specificity', cutoff)['estimate']:.3f}",
            ci(metric(pog, "primary_breast", "fixed_cutoff", "auroc", cutoff)),
        ]
        for cell, value in zip(cells, vals):
            cell.text = value
    set_table_geometry(cutoff_table, [680, 1050, 2350, 1200, 1200, 2880])
    format_table(cutoff_table, font_size=7.6)

    add_rich_paragraph(
        doc,
        [
            ("Calibration note: ", True, INK, False),
            (f"observed mean HRD score {observed_mean:.2f} vs predicted mean {predicted_mean:.2f}; breast R2 = {breast_r2:.3f}. Correlation was preserved, but absolute calibration was poor.", False, INK, False),
        ],
        after=2,
        size=8.7,
    )

    add_heading(doc, "2. PCAWG ICGC-only cancer-type validation")
    add_rich_paragraph(
        doc,
        [
            ("Cohort: ", True, INK, False),
            (f"{pcawg_usable}/{len(pcawg_mapping)} usable across {mapping_count} mappings; {pcawg_excluded} sample-level exclusions. Breast-DCIS and combined/Pancan files were excluded a priori. ", False, INK, False),
            ("Model: ", True, INK, False),
            (f"frozen 20-class standard_sbs96_id83 model ({cancer_feature_count} features).", False, INK, False),
        ],
        after=3,
    )
    summary_table = doc.add_table(rows=2, cols=4)
    summary_table.style = "Table Grid"
    metric_names = ["Overall accuracy", "Balanced accuracy", "Macro-F1 (mapped 9)", "Top-3 accuracy"]
    metric_keys = ["overall_accuracy", "balanced_accuracy", "macro_f1_mapped9", "top3_accuracy"]
    for index, (name, key) in enumerate(zip(metric_names, metric_keys)):
        summary_table.rows[0].cells[index].text = name
        row = pcawg_metric(pcawg, key)
        summary_table.rows[1].cells[index].text = ci(row)
    set_table_geometry(summary_table, [2340, 2340, 2340, 2340])
    format_table(summary_table, font_size=7.8, first_col_left=False)

    recalls = pcawg[pcawg["metric"].eq("per_cancer_recall")].copy()
    recall_text = ";  ".join(
        f"{row.class_label} {row.estimate:.3f} ({row.ci_low:.3f}-{row.ci_high:.3f})"
        for row in recalls.itertuples(index=False)
    )
    add_rich_paragraph(
        doc,
        [("Per-cancer recall (95% CI): ", True, INK, False), (recall_text, False, INK, False)],
        after=3,
        size=8.15,
    )

    add_heading(doc, "Interpretation and reproducibility")
    add_rich_paragraph(
        doc,
        [
            ("Cohort shift. ", True, INK, False),
            ("POG570/PCAWG are WGS cohorts; TCGA MC3 uses exome mutation calls. POG570 also contains advanced metastatic/recurrent cancers, whereas TCGA is predominantly primary tumors. ", False, INK, False),
            ("Model provenance. ", True, INK, False),
            ("Vijay's pushed files are new deployment artifacts selected by TCGA-only CV and refit on all eligible TCGA samples; they are not exports of discarded nested-CV fold models.", False, INK, False),
        ],
        after=2,
        size=8.6,
    )
    add_rich_paragraph(
        doc,
        [
            ("Reproducibility package: ", True, MUTED, False),
            ("sample mappings, exclusions, feature matrices, predictions, metrics, confusion matrix, input/model checksums, preprocessing, and software versions are saved under results/external_validation/aaron_email/.", False, MUTED, False),
        ],
        after=0,
        size=7.8,
    )

    doc.core_properties.title = "External Validation Results: POG570 and PCAWG"
    doc.core_properties.subject = "One-page external validation summary for Aaron Ge"
    doc.core_properties.author = "Cindy Liu"
    doc.core_properties.keywords = "POG570, HRD, PCAWG, external validation"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
